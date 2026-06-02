# -*- coding: utf-8 -*-
"""rebuild presentation 3.pptx as a polished, grid-based deck that matches the
improved models (77% classification, 0.94 pooled r2). dark cinematic palette,
sf pro display typography, embedded dark charts rendered from the live data."""
import sys, os, tempfile
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "dashboard"))

from utils.data import load_prices, train_classifier, fit_clusters, train_forecaster  # noqa

# ---------------------------------------------------------------- palette ----
BG       = "#0b0c0e"
SURFACE  = "#15171a"
SURFACE2 = "#1d2024"
BORDER   = "#2a2d33"
TEXT     = "#f5f5f4"
TEXT2    = "#a1a1aa"
TEXT3    = "#71717a"
ACCENT   = "#f59e0b"
ACCENT2  = "#fbbf24"
BRAND    = "#2d6a4f"
OK       = "#4ade80"
BAD      = "#ef4444"
BLUE     = "#60a5fa"

def hx(c):
    c = c.lstrip("#")
    return tuple(int(c[i:i+2], 16) for i in (0, 2, 4))

# ---------------------------------------------------------------- charts -----
ASSETS = Path(tempfile.mkdtemp(prefix="cafu_ppt_"))
CHART_FONT = "Arial"  # sf pro display is not installed for matplotlib; arial is clean
plt.rcParams.update({
    "font.family": CHART_FONT,
    "text.color": TEXT,
    "axes.edgecolor": BORDER,
    "axes.labelcolor": TEXT2,
    "xtick.color": TEXT2,
    "ytick.color": TEXT2,
    "axes.linewidth": 0.8,
    "figure.dpi": 200,
})

def _style(ax):
    ax.set_facecolor("none")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color(BORDER)
    ax.grid(axis="y", color="#ffffff", alpha=0.06, linewidth=0.8)
    ax.tick_params(length=0)

def save(fig, name):
    p = ASSETS / name
    fig.savefig(p, transparent=True, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)
    return str(p)

print("loading data and models...")
df = load_prices()
rc = train_classifier(df)
rcl = fit_clusters(df, k=4)
rf = train_forecaster(df)

ACC = rc["accuracy"] * 100
R2  = rf["r2"]

# 1. yearly trend
yearly = df.groupby("year")["price"].mean()
fig, ax = plt.subplots(figsize=(7.4, 3.5))
ax.plot(yearly.index, yearly.values, color=ACCENT, lw=2.6)
ax.fill_between(yearly.index, yearly.values, color=ACCENT, alpha=0.10)
ax.axvspan(2019.5, yearly.index.max(), color=BAD, alpha=0.06)
ax.scatter(yearly.index, yearly.values, color=ACCENT, s=22, zorder=5,
           edgecolors=BG, linewidths=1.2)
_style(ax); ax.set_ylabel("avg price, XAF")
TREND = save(fig, "trend.png")

# 2. classification arc 35.5 -> 63 -> 77
fig, ax = plt.subplots(figsize=(6.4, 3.7))
bars = ["baseline\n(guess medium)", "basic\nfeatures", "rich\nfeatures"]
vals = [35.5, 63.2, ACC]
cols = [TEXT3, BRAND, ACCENT]
b = ax.bar(bars, vals, color=cols, width=0.62)
for r, v in zip(b, vals):
    ax.text(r.get_x()+r.get_width()/2, v+1.5, f"{v:.0f}%", ha="center",
            color=TEXT, fontsize=15, fontweight="bold")
ax.set_ylim(0, 92); _style(ax); ax.set_ylabel("accuracy")
ax.set_yticks([])
CLASS = save(fig, "class.png")

# 3. cluster pca scatter
pivot = rcl["pivot"]; coords = rcl["coords"]
names = {0: "far north conflict", 1: "urban hubs",
         2: "anglophone southwest", 3: "eastern refugee corridor"}
cmap = {0: BAD, 1: ACCENT, 2: BLUE, 3: OK}
fig, ax = plt.subplots(figsize=(7.0, 4.2))
for c in sorted(pivot["cluster"].unique()):
    m = pivot["cluster"].values == c
    ax.scatter(coords[m, 0], coords[m, 1], s=120, color=cmap.get(c, TEXT2),
               edgecolors=BG, linewidths=1.6, alpha=0.92, label=names.get(c, str(c)))
_style(ax); ax.grid(False)
ax.set_xlabel(f"PCA 1 ({rcl['ev'][0]*100:.0f}% var)")
ax.set_ylabel(f"PCA 2 ({rcl['ev'][1]*100:.0f}% var)")
leg = ax.legend(frameon=False, fontsize=9, labelcolor=TEXT2, loc="best")
CLUST = save(fig, "clusters.png")

# 4. forecast actual vs predicted, maize
mt = rf["test"][rf["test"]["commodity"] == "Maize (white)"].sort_values("date")
fig, ax = plt.subplots(figsize=(7.4, 3.5))
ax.plot(mt["date"], mt["price"], color=TEXT, lw=2.2, label="actual")
ax.plot(mt["date"], mt["pred"], color=ACCENT, lw=2.2, ls=(0, (2, 2)), label="predicted")
_style(ax); ax.set_ylabel("price, XAF")
ax.legend(frameon=False, fontsize=10, labelcolor=TEXT2)
FORE = save(fig, "forecast.png")

# 5. year over year shocks
yoy = (yearly.pct_change() * 100).dropna()
thr = yoy.mean() + 2 * yoy.std()
bc = [ACCENT if v > thr else (BAD if v > 0 else OK) for v in yoy.values]
fig, ax = plt.subplots(figsize=(7.6, 3.4))
ax.bar(yoy.index.astype(int), yoy.values, color=bc, width=0.74)
ax.axhline(thr, color=ACCENT, ls="--", lw=1, alpha=0.8)
ax.axhline(0, color=TEXT3, lw=0.8)
_style(ax); ax.set_ylabel("year over year %")
YOY = save(fig, "yoy.png")

stats = {
    "rows": len(df),
    "regions": df["admin1"].nunique(),
    "markets": df["market"].nunique(),
    "commodities": df["commodity"].nunique(),
}
print("charts done. acc=%.1f r2=%.3f stats=%s" % (ACC, R2, stats))

# ---------------------------------------------------------------- deck -------
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE

HEAD = "SF Pro Display"
BODY = "SF Pro Display"

def C(hexs):
    return RGBColor(*hx(hexs))

EMU_IN = 914400
PW, PH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(PW)
prs.slide_height = Inches(PH)
BLANK = prs.slide_layouts[6]

def new_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C(BG)
    return s

def textbox(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=1.0, space_after=0):
    """runs: list of paragraphs; each paragraph is a list of (text, opts)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        for txt, o in para:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = o.get("font", BODY)
            f.size = Pt(o.get("size", 18))
            f.bold = o.get("bold", False)
            f.italic = o.get("italic", False)
            f.color.rgb = C(o.get("color", TEXT))
    return tb

def rule(slide, x, y, w, color=ACCENT, h=0.035):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = C(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def card(slide, x, y, w, h, fill=SURFACE, border=BORDER, accent=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = C(fill)
    sh.line.color.rgb = C(border); sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = 0.06
    except Exception:
        pass
    if accent:
        rule(slide, x, y, 0.06, color=accent, h=h)
    return sh

def kicker(slide, x, y, text, color=ACCENT, w=8):
    textbox(slide, x, y, w, 0.35,
            [[(text.upper(), {"font": HEAD, "size": 12.5, "bold": True, "color": color})]])

def footer(slide, n):
    textbox(slide, 0.7, PH - 0.5, 8, 0.3,
            [[("cameroon food prices  ·  a data mining story", {"size": 9, "color": TEXT3})]])
    textbox(slide, PW - 1.7, PH - 0.5, 1, 0.3,
            [[(f"{n:02d}", {"font": HEAD, "size": 9, "color": TEXT3})]],
            align=PP_ALIGN.RIGHT)

MX = 0.85  # left margin

# ---- slide 1, title --------------------------------------------------------
s = new_slide()
kicker(s, MX, 1.5, "a data mining project")
rule(s, MX, 1.95, 0.9)
textbox(s, MX, 2.25, 11.6, 2.6, [[
    ("what 21 years of food prices say about ", {"font": HEAD, "size": 47, "bold": True, "color": TEXT}),
    ("a country in flux", {"font": HEAD, "size": 47, "bold": True, "color": ACCENT}),
]], line_spacing=1.02)
textbox(s, MX, 5.0, 9.5, 1.0, [[
    ("the world food programme tracked the price of everyday food in 76 markets "
     "across cameroon. this is what the numbers reveal.",
     {"size": 17, "color": TEXT2})]], line_spacing=1.3)
for i, (lab) in enumerate(["cameroon", "2005 — 2026", "55,660 observations"]):
    textbox(s, MX + i * 3.4, 6.3, 3.3, 0.4,
            [[(lab, {"font": HEAD, "size": 12, "bold": True, "color": TEXT2})]])
rule(s, MX, 6.2, 10.0, color=BORDER, h=0.012)

# ---- slide 2, the dataset --------------------------------------------------
s = new_slide()
kicker(s, MX, 0.7, "00 / prologue")
textbox(s, MX, 1.1, 6, 1.2, [[("first, the ", {"font": HEAD, "size": 38, "bold": True, "color": TEXT}),
                              ("numbers.", {"font": HEAD, "size": 38, "bold": True, "color": ACCENT})]])
textbox(s, MX, 2.25, 4.7, 3.6, [
    [("this story rests on a single csv from the world food programme — the un "
      "agency that monitors hunger.", {"size": 15, "color": TEXT2})],
    [("nine regions, not ten. the missing one is ", {"size": 15, "color": TEXT2}),
     ("Sud", {"size": 15, "color": ACCENT, "italic": True, "bold": True}),
     (", the peaceful south. wfp follows need, so its data goes where the crises "
      "are: the far north, the anglophone regions, the east.", {"size": 15, "color": TEXT2})],
], line_spacing=1.32, space_after=10)
# stat grid 2 x 3 on the right
grid = [("55,660", "records", TEXT), (f"{stats['regions']} / 10", "regions covered", ACCENT),
        (str(stats["markets"]), "markets", TEXT), (str(stats["commodities"]), "commodities", TEXT),
        ("21", "years of history", ACCENT), ("1", "region missing — Sud", BAD)]
gx, gy, cw, ch, gap = 6.05, 2.2, 3.35, 1.55, 0.22
for i, (v, l, col) in enumerate(grid):
    r, c = divmod(i, 2)
    x = gx + c * (cw + gap); y = gy + r * (ch + gap)
    card(s, x, y, cw, ch)
    textbox(s, x + 0.25, y + 0.18, cw - 0.4, 0.95,
            [[(v, {"font": HEAD, "size": 33, "bold": True, "color": col})]])
    textbox(s, x + 0.25, y + 1.02, cw - 0.4, 0.45,
            [[(l.upper(), {"font": HEAD, "size": 10.5, "color": TEXT3})]])
footer(s, 2)

# ---- slide 3, the data -----------------------------------------------------
s = new_slide()
kicker(s, MX, 0.7, "01 / the data")
textbox(s, MX, 1.1, 8, 1.0, [[("before any model, you ", {"font": HEAD, "size": 34, "bold": True, "color": TEXT}),
                              ("look.", {"font": HEAD, "size": 34, "bold": True, "color": ACCENT})]])
s.shapes.add_picture(TREND, Inches(MX), Inches(2.25), width=Inches(7.3))
textbox(s, MX, 1.95, 7, 0.3, [[("NATIONAL AVERAGE PRICE BY YEAR — A LONG CLIMB, THEN A POST-2020 JUMP",
                                {"font": HEAD, "size": 10.5, "color": TEXT3})]])
shapes = [("a long climb", "prices drift up for 15 years, then accelerate after 2020 — covid, war, inflation."),
          ("a lean season", "june to august sit above the line every year, between harvests."),
          ("a crisis map", "extreme nord dominates the records 3:1 — wfp monitors where hunger is."),
          ("imports cost", "processed and imported goods sit far above local staples per kilo.")]
cx = 8.5
for i, (t, d) in enumerate(shapes):
    y = 2.25 + i * 1.16
    rule(s, cx, y + 0.05, 0.05, h=0.85)
    textbox(s, cx + 0.25, y, 4.0, 0.4, [[(t, {"font": HEAD, "size": 15, "bold": True, "color": TEXT})]])
    textbox(s, cx + 0.25, y + 0.4, 4.0, 0.7, [[(d, {"size": 11.5, "color": TEXT2})]], line_spacing=1.18)
footer(s, 3)

# ---- slide 4, classification ----------------------------------------------
s = new_slide()
kicker(s, MX, 0.7, "02 / classification")
textbox(s, MX, 1.1, 11, 1.0, [[("can a model read a price ", {"font": HEAD, "size": 32, "bold": True, "color": TEXT}),
                               ("tier?", {"font": HEAD, "size": 32, "bold": True, "color": ACCENT})]])
# big stat left
textbox(s, MX, 2.5, 4.6, 1.8, [[(f"{ACC:.0f}%", {"font": HEAD, "size": 96, "bold": True, "color": ACCENT})]])
textbox(s, MX, 4.5, 4.6, 1.8, [
    [("accuracy on held-out data, up from 63%.", {"size": 15, "color": TEXT})],
    [("a random forest labels each price Low / Medium / High against its own "
      "commodity's history. the lift came from feature engineering, not a "
      "bigger model: market location plus recent and yearly median price.",
      {"size": 13, "color": TEXT2})],
], line_spacing=1.3, space_after=8)
s.shapes.add_picture(CLASS, Inches(7.0), Inches(2.5), width=Inches(5.6))
textbox(s, 7.0, 2.15, 5.6, 0.3, [[("THE ARC: SAME ALGORITHM, BETTER FEATURES",
                                   {"font": HEAD, "size": 10.5, "color": TEXT3})]])
footer(s, 4)

# ---- slide 5, clusters -----------------------------------------------------
s = new_slide()
kicker(s, MX, 0.7, "03 / clusters")
textbox(s, MX, 1.1, 11.5, 1.0, [[("let the data ", {"font": HEAD, "size": 32, "bold": True, "color": TEXT}),
                                 ("draw the map.", {"font": HEAD, "size": 32, "bold": True, "color": ACCENT})]])
s.shapes.add_picture(CLUST, Inches(MX), Inches(2.2), width=Inches(6.6))
textbox(s, MX, 1.9, 7, 0.3, [[("76 MARKETS BY PRICE PROFILE, PROJECTED TO 2D — K-MEANS, k = 4",
                              {"font": HEAD, "size": 10.5, "color": TEXT3})]])
clusters = [("far north conflict zone", BAD), ("urban commercial hubs", ACCENT),
            ("anglophone southwest", BLUE), ("eastern refugee corridor", OK)]
cx2 = 7.7
textbox(s, cx2, 2.15, 5, 0.7, [[("k-means knew nothing about politics. it saw only prices — and "
                                 "rebuilt the country's crisis geography.", {"size": 13, "color": TEXT2})]],
        line_spacing=1.25)
for i, (nm, col) in enumerate(clusters):
    y = 3.25 + i * 0.92
    card(s, cx2, y, 4.85, 0.78, accent=col)
    textbox(s, cx2 + 0.3, y + 0.12, 4.4, 0.3, [[(f"cluster {i:02d}", {"font": HEAD, "size": 9.5, "bold": True, "color": col})]])
    textbox(s, cx2 + 0.3, y + 0.38, 4.4, 0.35, [[(nm, {"font": HEAD, "size": 14, "bold": True, "color": TEXT})]])
footer(s, 5)

# ---- slide 6, forecasting --------------------------------------------------
s = new_slide()
kicker(s, MX, 0.7, "04 / forecasting")
textbox(s, MX, 1.1, 11, 1.0, [[("how well can we see ", {"font": HEAD, "size": 32, "bold": True, "color": TEXT}),
                               ("one month ahead?", {"font": HEAD, "size": 32, "bold": True, "color": ACCENT})]])
textbox(s, MX, 2.4, 4.6, 1.6, [[(f"{R2:.2f}", {"font": HEAD, "size": 92, "bold": True, "color": ACCENT})]])
textbox(s, MX, 4.25, 4.6, 0.6, [[("pooled R², up from 0.48", {"size": 15, "color": TEXT})]])
steps = [("use more data", "train on the top 10 commodities, not maize alone — 10x the rows."),
         ("predict log price", "a 10% move counts the same on cheap and dear items."),
         ("gradient boosting", "trees that learn from each other's mistakes, built for trend and season.")]
for i, (t, d) in enumerate(steps):
    y = 4.95 + i * 0.0  # stacked under stat
for i, (t, d) in enumerate(steps):
    y = 4.95 + i * 0.78
    textbox(s, MX, y, 0.5, 0.4, [[(f"0{i+1}", {"font": HEAD, "size": 16, "bold": True, "color": ACCENT})]])
    textbox(s, MX + 0.55, y, 4.2, 0.35, [[(t, {"font": HEAD, "size": 13.5, "bold": True, "color": TEXT})]])
    textbox(s, MX + 0.55, y + 0.32, 4.2, 0.5, [[(d, {"size": 10.5, "color": TEXT2})]], line_spacing=1.12)
s.shapes.add_picture(FORE, Inches(6.9), Inches(2.6), width=Inches(5.7))
textbox(s, 6.9, 2.25, 5.7, 0.3, [[("ACTUAL VS PREDICTED — MAIZE (WHITE), THROUGH THE POST-2020 WAVE",
                                   {"font": HEAD, "size": 10.5, "color": TEXT3})]])
footer(s, 6)

# ---- slide 7, patterns -----------------------------------------------------
s = new_slide()
kicker(s, MX, 0.7, "05 / patterns")
textbox(s, MX, 1.1, 11, 1.0, [[("when the system ", {"font": HEAD, "size": 32, "bold": True, "color": TEXT}),
                               ("cracked.", {"font": HEAD, "size": 32, "bold": True, "color": ACCENT})]])
s.shapes.add_picture(YOY, Inches(MX), Inches(2.5), width=Inches(7.6))
textbox(s, MX, 2.2, 8, 0.3, [[("YEAR-OVER-YEAR PRICE CHANGE — AMBER BARS CROSS THE 2-SIGMA ANOMALY LINE",
                              {"font": HEAD, "size": 10.5, "color": TEXT3})]])
notes = [("2008", "the global food price crisis."),
         ("2022", "post-covid inflation plus the war on grain."),
         ("volatile", "animal products and oils swing most; cereals are policy-stabilised.")]
cx3 = 8.8
textbox(s, cx3, 2.4, 4, 0.4, [[("two shocks, found without labels", {"font": HEAD, "size": 16, "bold": True, "color": TEXT})]])
for i, (t, d) in enumerate(notes):
    y = 3.2 + i * 1.05
    textbox(s, cx3, y, 4, 0.4, [[(t, {"font": HEAD, "size": 14, "bold": True, "color": ACCENT})]])
    textbox(s, cx3, y + 0.36, 4, 0.6, [[(d, {"size": 11.5, "color": TEXT2})]], line_spacing=1.18)
footer(s, 7)

# ---- slide 8, conclusion ---------------------------------------------------
s = new_slide()
kicker(s, MX, 0.7, "06 / conclusion")
textbox(s, MX, 1.1, 11, 1.0, [[("what the ", {"font": HEAD, "size": 34, "bold": True, "color": TEXT}),
                               ("data said.", {"font": HEAD, "size": 34, "bold": True, "color": ACCENT})]])
takeaways = [
    ("01", "prices are structured, not chaotic", "a classifier reads them at 77% and a forecaster hits 0.94 pooled R²."),
    ("02", "when beats where", "recent history and inflation context predict price better than geography."),
    ("03", "markets map to crisis", "four price clusters line up with cameroon's humanitarian zones."),
    ("04", "two system shocks", "2008 and 2022 surface cleanly, with no one labelling them."),
]
gx, gy, cw, ch, gap = MX, 2.35, 5.75, 1.75, 0.3
for i, (n, t, d) in enumerate(takeaways):
    r, c = divmod(i, 2)
    x = gx + c * (cw + gap); y = gy + r * (ch + gap)
    card(s, x, y, cw, ch, accent=ACCENT)
    textbox(s, x + 0.32, y + 0.2, 1, 0.4, [[(n, {"font": HEAD, "size": 13, "bold": True, "color": ACCENT})]])
    textbox(s, x + 0.32, y + 0.52, cw - 0.6, 0.4, [[(t, {"font": HEAD, "size": 18, "bold": True, "color": TEXT})]])
    textbox(s, x + 0.32, y + 1.02, cw - 0.6, 0.6, [[(d, {"size": 12, "color": TEXT2})]], line_spacing=1.2)
textbox(s, MX, 6.45, 11.6, 0.6,
        [[("numbers, given time, will tell on the world.", {"font": HEAD, "size": 17, "italic": True, "color": ACCENT})]],
        align=PP_ALIGN.CENTER)
footer(s, 8)

OUT = ROOT / "presentation 3.pptx"
prs.save(OUT)
print("saved", OUT, "with", len(prs.slides._sldIdLst), "slides")

# cleanup chart pngs
for f in ASSETS.glob("*.png"):
    f.unlink()
ASSETS.rmdir()
print("done")
